from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required. Install it with: python -m pip install openpyxl"
    ) from exc


SCAN_ROWS = 900
SCAN_COLS = 90
SUPPLIER_AVANSUM = "Avansum"
SUPPLIER_OS_PLUS = "OS+ materials"


@dataclass(frozen=True)
class MatrixCell:
    width_mm: int
    height_mm: int
    price: float


@dataclass(frozen=True)
class PriceMatrix:
    supplier: str
    name: str
    profile_system: str
    glass_unit: str
    color: str
    matrix_kind: str
    hardware_type: str
    variant_label: str
    source_file: str
    source_sheet: str
    source_block: int
    effective_from: str | None
    meta: dict[str, Any]
    cells: list[MatrixCell]


@dataclass(frozen=True)
class CatalogItem:
    supplier: str
    name: str
    sku: str
    category: str
    unit: str
    buy_price: float
    sell_price: float
    item_type: str
    source_file: str
    source_sheet: str
    notes: str = ""


@dataclass(frozen=True)
class WorkRate:
    name: str
    category: str
    unit: str
    rate: float
    source_file: str


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean(value)
    if not text:
        return None
    text = text.replace("≈", "").replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def int_number(value: Any) -> int | None:
    parsed = number(value)
    if parsed is None:
        return None
    if abs(parsed - round(parsed)) > 0.001:
        return None
    return int(round(parsed))


def date_value(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def sql_text(value: Any) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def sql_num(value: float | int | None) -> str:
    if value is None:
        return "null"
    return f"{float(value):.2f}".rstrip("0").rstrip(".")


def sql_json(value: Any) -> str:
    return sql_text(json.dumps(value, ensure_ascii=False, sort_keys=True)) + "::jsonb"


def sku(prefix: str, *parts: str) -> str:
    raw = "|".join(clean(part).lower() for part in parts if clean(part))
    digest = hashlib.md5(raw.encode("utf-8")).hexdigest()[:12].upper()
    return f"{prefix}-{digest}"


def safe_label(*parts: str) -> str:
    raw = "-".join(clean(part).lower() for part in parts if clean(part))
    raw = re.sub(r"[^a-zа-я0-9]+", "-", raw, flags=re.IGNORECASE).strip("-")
    return raw[:80] or "standard"


def iter_rows(ws, max_rows: int = SCAN_ROWS, max_cols: int = SCAN_COLS) -> list[list[Any]]:
    return [
        list(row)
        for row in ws.iter_rows(
            min_row=1,
            max_row=max_rows,
            min_col=1,
            max_col=max_cols,
            values_only=True,
        )
    ]


def cell(row: list[Any], zero_col: int) -> Any:
    return row[zero_col] if 0 <= zero_col < len(row) else None


def find_file(default_dir: Path, exact_name: str, startswith: str | None = None) -> Path:
    exact = default_dir / exact_name
    if exact.exists():
        return exact
    if startswith:
        matches = sorted(default_dir.glob(f"{startswith}*.xlsx"))
        if matches:
            return matches[0]
    raise FileNotFoundError(f"Workbook not found: {exact}")


def infer_profile_and_glass(sheet_name: str, profile_text: str, filling: str) -> tuple[str, str]:
    source = f"{profile_text} {sheet_name}"
    glass_match = re.search(r"СП\s*\d+|\bSP\s*\d+", source, flags=re.IGNORECASE)
    if glass_match:
        glass_unit = glass_match.group(0).replace(" ", "").upper()
    else:
        fill_match = re.search(r"\d+\s*[МM]1(?:[-–]\d+[-–]\d+\s*[МM]1)+|\d+\s*[МM]1[-–]\d+[-–]\d+\s*[МM]1", filling)
        glass_unit = clean(fill_match.group(0)) if fill_match else clean(filling)[:48] or "glass_unit"

    profile = clean(profile_text) or sheet_name
    profile = re.sub(r"\(?\s*СП\s*\d+\s*\)?", "", profile, flags=re.IGNORECASE)
    profile = re.sub(r"\(?\s*SP\s*\d+\s*\)?", "", profile, flags=re.IGNORECASE)
    return clean(profile), glass_unit


def infer_kind(sheet_name: str, profile_text: str) -> str:
    haystack = f"{sheet_name} {profile_text}".lower()
    if "двер" in haystack:
        return "door"
    if "балкон" in haystack:
        return "balcony"
    return "window"


def parse_avansum_matrices(path: Path) -> list[PriceMatrix]:
    wb = load_workbook(path, read_only=True, data_only=True)
    matrices: list[PriceMatrix] = []
    for ws in wb.worksheets:
        rows = iter_rows(ws)
        effective_from = next(
            (date_value(value) for row in rows[:8] for value in row if date_value(value)),
            None,
        )
        block_no = 0
        for r_idx, row in enumerate(rows):
            marker_col = next(
                (
                    c_idx
                    for c_idx, value in enumerate(row)
                    if clean(value).lower().startswith("профиль:")
                ),
                None,
            )
            if marker_col is None:
                continue
            header_col = next(
                (
                    c_idx
                    for c_idx, value in enumerate(row)
                    if clean(value).replace("/", "\\").upper() in {"H\\B", "Н\\В"}
                ),
                None,
            )
            if header_col is None:
                continue

            widths: list[tuple[int, int]] = []
            for c_idx in range(header_col + 1, len(row)):
                width = int_number(row[c_idx])
                if width is None:
                    if widths:
                        break
                    continue
                widths.append((c_idx, width))
            if not widths:
                continue

            block_no += 1
            profile_text = clean(cell(row, marker_col + 1))
            color = clean(cell(rows[r_idx + 1], marker_col + 1)) if r_idx + 1 < len(rows) else ""
            hardware = clean(cell(rows[r_idx + 2], marker_col + 1)) if r_idx + 2 < len(rows) else ""
            filling = clean(cell(rows[r_idx + 3], marker_col + 1)) if r_idx + 3 < len(rows) else ""
            profile_system, glass_unit = infer_profile_and_glass(ws.title, profile_text, filling)
            matrix_kind = infer_kind(ws.title, profile_text)

            cells: list[MatrixCell] = []
            heights: list[int] = []
            for data_row in rows[r_idx + 1 :]:
                height = int_number(cell(data_row, header_col))
                if height is None:
                    if heights:
                        break
                    continue
                heights.append(height)
                for c_idx, width in widths:
                    price = number(cell(data_row, c_idx))
                    if price is not None and price > 0:
                        cells.append(MatrixCell(width_mm=width, height_mm=height, price=price))
            if not cells:
                continue

            variant = safe_label(str(block_no), matrix_kind, hardware, filling)
            name = clean(f"{ws.title}: {profile_system}, {glass_unit}, {color or 'Белый'}, block {block_no}")
            matrices.append(
                PriceMatrix(
                    supplier=SUPPLIER_AVANSUM,
                    name=name,
                    profile_system=profile_system,
                    glass_unit=glass_unit,
                    color=color or "Белый",
                    matrix_kind=matrix_kind,
                    hardware_type=hardware or "standard",
                    variant_label=variant,
                    source_file=path.name,
                    source_sheet=ws.title,
                    source_block=block_no,
                    effective_from=effective_from,
                    meta={
                        "filling": filling,
                        "header_row": r_idx + 1,
                        "widths": [width for _, width in widths],
                        "heights": heights,
                    },
                    cells=cells,
                )
            )
    return matrices


def infer_unit(name: str) -> str:
    lower = name.lower()
    if "м2" in lower or "м²" in lower or "м.кв" in lower:
        return "м2"
    if "м пог" in lower or "м.пог" in lower or "м. пог" in lower:
        return "м.пог"
    if "комп" in lower:
        return "комплект"
    if "кг" in lower:
        return "кг"
    if "л " in lower or lower.endswith(" л"):
        return "л"
    if "рулон" in lower:
        return "рулон"
    return "шт"


def infer_category(name: str, fallback: str = "materials") -> str:
    lower = name.lower()
    pairs = [
        ("hardware", ["фурнит", "ручка", "петл", "замок", "гарнитур"]),
        ("profiles", ["профил", "расширител", "соединител", "нащельник"]),
        ("glass", ["стекл", "тонир"]),
        ("floor", ["пол", "osb", "фанера", "ламинат", "линолеум"]),
        ("insulation", ["утепл", "пенопл", "техноплекс", "пенофол", "вата"]),
        ("panels", ["панел", "сайдинг", "вагонка", "гипс"]),
        ("fasteners", ["саморез", "дюб", "анк", "креп"]),
        ("consumables", ["пена", "гермет", "клей", "скотч", "лента"]),
        ("finishing", ["плинтус", "откос", "подокон", "штукатур"]),
    ]
    return next((category for category, keys in pairs if any(key in lower for key in keys)), fallback)


def parse_os_materials(path: Path) -> list[CatalogItem]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    items: list[CatalogItem] = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row or 1200, min_col=1, max_col=4, values_only=True):
        name = clean(row[0])
        if not name:
            continue
        buy_price = number(row[1]) or 0
        coeff = number(row[2])
        sell_price = number(row[3])
        if sell_price is None and coeff and buy_price:
            sell_price = buy_price * coeff
        sell_price = sell_price or buy_price
        unit = infer_unit(name)
        category = infer_category(name)
        items.append(
            CatalogItem(
                supplier=SUPPLIER_OS_PLUS,
                name=name,
                sku=sku("OSM", name, unit),
                category=category,
                unit=unit,
                buy_price=buy_price,
                sell_price=sell_price,
                item_type="material",
                source_file=path.name,
                source_sheet=ws.title,
                notes=f"coefficient={coeff}" if coeff else "",
            )
        )
    return items


def parse_os_work_rates(path: Path) -> list[WorkRate]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    rates: list[WorkRate] = []
    current_name = ""
    current_unit = ""
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row or 1200, min_col=1, max_col=4, values_only=True):
        if clean(row[0]):
            current_name = clean(row[0])
        if clean(row[2]):
            current_unit = clean(row[2])
        category = clean(row[1])
        rate = number(row[3])
        if not current_name or rate is None or rate <= 0:
            continue
        rates.append(
            WorkRate(
                name=current_name,
                category=category,
                unit=current_unit or "шт",
                rate=rate,
                source_file=path.name,
            )
        )
    return rates


def parse_avansum_catalog(path: Path) -> list[CatalogItem]:
    wb = load_workbook(path, read_only=True, data_only=True)
    items: list[CatalogItem] = []
    list_sheets = [ws for ws in wb.worksheets if ws.title in {
        "Фурнитура",
        "Фурнитура оконная",
        "Фурнитура дверная",
        "Доп. профили",
    }]
    for ws in list_sheets:
        rows = iter_rows(ws, max_rows=300, max_cols=16)
        for row in rows[7:]:
            candidates = [clean(value) for value in row[:8] if clean(value)]
            text_candidates = [value for value in candidates if not re.fullmatch(r"-?\d+(?:[.,]\d+)?", value)]
            if not text_candidates:
                continue
            price_candidates = [number(value) for value in row[:8]]
            prices = [value for value in price_candidates if value is not None and value > 0]
            if not prices:
                continue
            name = max(text_candidates, key=len)
            if len(name) < 5 or "прайс" in name.lower():
                continue
            color = clean(row[4]) if len(row) > 4 else ""
            unit = infer_unit(name)
            category = infer_category(f"{ws.title} {name}", fallback="window_components")
            items.append(
                CatalogItem(
                    supplier=SUPPLIER_AVANSUM,
                    name=f"{name} ({color})" if color and color.lower() not in name.lower() else name,
                    sku=sku("AVA", ws.title, name, color, unit),
                    category=category,
                    unit=unit,
                    buy_price=prices[-1],
                    sell_price=prices[-1],
                    item_type="component",
                    source_file=path.name,
                    source_sheet=ws.title,
                )
            )
    dedup: dict[str, CatalogItem] = {}
    for item in items:
        dedup[item.sku] = item
    return list(dedup.values())


def supplier_sql(name: str, categories: list[str], notes: str) -> str:
    return f"""
insert into public.suppliers (name, categories, notes, is_active)
values ({sql_text(name)}, array[{", ".join(sql_text(c) for c in categories)}]::text[], {sql_text(notes)}, true)
on conflict (name) do update set
  categories = excluded.categories,
  notes = excluded.notes,
  is_active = true,
  updated_at = now();
""".strip()


def catalog_item_sql(item: CatalogItem) -> str:
    return f"""
insert into public.warehouse_items (name, category, unit, sku, item_type, sell_price, notes, is_active)
values (
  {sql_text(item.name)}, {sql_text(item.category)}, {sql_text(item.unit)}, {sql_text(item.sku)},
  {sql_text(item.item_type)}, {sql_num(item.sell_price)}, {sql_text(f"{item.source_file} / {item.source_sheet}. {item.notes}".strip())}, true
)
on conflict (sku) do update set
  name = excluded.name,
  category = excluded.category,
  unit = excluded.unit,
  item_type = excluded.item_type,
  sell_price = excluded.sell_price,
  notes = excluded.notes,
  is_active = true,
  updated_at = now();

with supplier as (
  select id from public.suppliers where name = {sql_text(item.supplier)}
)
insert into public.supplier_items (supplier_id, name, sku, unit, category, buy_price, sell_price, availability, is_active)
select supplier.id, {sql_text(item.name)}, {sql_text(item.sku)}, {sql_text(item.unit)}, {sql_text(item.category)},
       {sql_num(item.buy_price)}, {sql_num(item.sell_price)}, 'order_only', true
from supplier
on conflict (supplier_id, sku) do update set
  name = excluded.name,
  unit = excluded.unit,
  category = excluded.category,
  buy_price = excluded.buy_price,
  sell_price = excluded.sell_price,
  is_active = true,
  updated_at = now();

insert into public.catalog_links (warehouse_item_id, supplier_item_id, is_primary, priority)
select wi.id, si.id, true, 10
from public.warehouse_items wi
join public.supplier_items si on si.sku = wi.sku
join public.suppliers s on s.id = si.supplier_id
where wi.sku = {sql_text(item.sku)} and s.name = {sql_text(item.supplier)}
on conflict (warehouse_item_id, supplier_item_id) do update set
  is_primary = excluded.is_primary,
  priority = excluded.priority;
""".strip()


def matrix_sql(matrix: PriceMatrix) -> str:
    values = ",\n".join(
        f"  ((select id from matrix), {cell.width_mm}, {cell.height_mm}, {sql_num(cell.price)})"
        for cell in matrix.cells
    )
    return f"""
with supplier as (
  select id from public.suppliers where name = {sql_text(matrix.supplier)}
),
upserted as (
  insert into public.window_price_matrices (
    supplier_id, name, profile_system, glass_unit, color, effective_from,
    source_file, source_sheet, matrix_kind, hardware_type, variant_label, source_block, meta,
    rounding_policy, is_active
  )
  select supplier.id, {sql_text(matrix.name)}, {sql_text(matrix.profile_system)}, {sql_text(matrix.glass_unit)},
         {sql_text(matrix.color)}, {sql_text(matrix.effective_from)}, {sql_text(matrix.source_file)},
         {sql_text(matrix.source_sheet)}, {sql_text(matrix.matrix_kind)}, {sql_text(matrix.hardware_type)},
         {sql_text(matrix.variant_label)}, {matrix.source_block}, {sql_json(matrix.meta)},
         'ceil_to_matrix', true
  from supplier
  on conflict (profile_system, glass_unit, color, source_sheet, source_block) do update set
    supplier_id = excluded.supplier_id,
    name = excluded.name,
    effective_from = excluded.effective_from,
    matrix_kind = excluded.matrix_kind,
    hardware_type = excluded.hardware_type,
    variant_label = excluded.variant_label,
    meta = excluded.meta,
    is_active = true,
    updated_at = now()
  returning id
),
matrix as (
  select id from upserted
  union all
  select id from public.window_price_matrices
  where profile_system = {sql_text(matrix.profile_system)}
    and glass_unit = {sql_text(matrix.glass_unit)}
    and color = {sql_text(matrix.color)}
    and source_sheet = {sql_text(matrix.source_sheet)}
    and source_block = {matrix.source_block}
  limit 1
)
insert into public.window_price_matrix_cells (matrix_id, width_mm, height_mm, price)
values
{values}
on conflict (matrix_id, width_mm, height_mm) do update set
  price = excluded.price;
""".strip()


def work_rate_sql(rate: WorkRate) -> str:
    return f"""
insert into public.window_work_rates (name, category, unit, rate, source_file, is_active)
values ({sql_text(rate.name)}, {sql_text(rate.category)}, {sql_text(rate.unit)}, {sql_num(rate.rate)}, {sql_text(rate.source_file)}, true)
on conflict (name, category, unit) do update set
  rate = excluded.rate,
  source_file = excluded.source_file,
  is_active = true,
  updated_at = now();
""".strip()


def write_sql(output: Path, matrices: list[PriceMatrix], items: list[CatalogItem], rates: list[WorkRate]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    chunks = [
        "-- Generated by scripts/import-window-workbooks.py",
        "-- Review before applying to production Supabase.",
        "begin;",
        supplier_sql(SUPPLIER_AVANSUM, ["windows", "doors", "profiles", "hardware"], "Imported from Avansum workbook."),
        supplier_sql(SUPPLIER_OS_PLUS, ["materials", "balcony_finishing", "labor"], "Imported from OS+ purchases workbook."),
    ]
    chunks.extend(catalog_item_sql(item) for item in items)
    chunks.extend(work_rate_sql(rate) for rate in rates)
    chunks.extend(matrix_sql(matrix) for matrix in matrices)
    chunks.append("commit;")
    output.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")


def write_preview(output: Path, matrices: list[PriceMatrix], items: list[CatalogItem], rates: list[WorkRate]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "counts": {
            "price_matrices": len(matrices),
            "price_matrix_cells": sum(len(matrix.cells) for matrix in matrices),
            "catalog_items": len(items),
            "work_rates": len(rates),
        },
        "samples": {
            "price_matrices": [
                {**asdict(matrix), "cells": [asdict(cell) for cell in matrix.cells[:5]]}
                for matrix in matrices[:5]
            ],
            "catalog_items": [asdict(item) for item in items[:12]],
            "work_rates": [asdict(rate) for rate in rates[:12]],
        },
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import window business workbooks into Supabase seed SQL.")
    default_dir = Path.home() / "Downloads" / "Telegram Desktop"
    parser.add_argument("--avansum", type=Path, default=None, help="Path to 69c38967982f6.xlsx")
    parser.add_argument("--os-plus", type=Path, default=None, help="Path to Закупки ОС+.xlsx")
    parser.add_argument("--out-sql", type=Path, default=Path("supabase/seed/window_import_seed.sql"))
    parser.add_argument("--out-preview", type=Path, default=Path("supabase/seed/window_import_preview.json"))
    args = parser.parse_args()

    avansum_path = args.avansum or find_file(default_dir, "69c38967982f6.xlsx", "69c38967982f6")
    os_plus_path = args.os_plus or find_file(default_dir, "Закупки ОС+.xlsx")

    matrices = parse_avansum_matrices(avansum_path)
    items = parse_avansum_catalog(avansum_path) + parse_os_materials(os_plus_path)
    rates = parse_os_work_rates(os_plus_path)

    write_sql(args.out_sql, matrices, items, rates)
    write_preview(args.out_preview, matrices, items, rates)

    print(json.dumps({
        "avansum": str(avansum_path),
        "os_plus": str(os_plus_path),
        "out_sql": str(args.out_sql),
        "out_preview": str(args.out_preview),
        "counts": {
            "price_matrices": len(matrices),
            "price_matrix_cells": sum(len(matrix.cells) for matrix in matrices),
            "catalog_items": len(items),
            "work_rates": len(rates),
        },
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
